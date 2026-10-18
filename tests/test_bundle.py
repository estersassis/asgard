"""Run manifest, file descriptions, and the optional classification spreadsheet."""

import builtins
import hashlib
import json
from importlib.metadata import version
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asgard.core.config import ConfigError, load_run
from asgard.core.prompt_registry import PromptRegistry
from asgard.core.stage import RunContext
from asgard.core.stages.bundle import BundleStage
from asgard.prompts.classifier import SYSTEM_PROMPT as CLASSIFIER_PROMPT
from asgard.prompts.evaluator import SYSTEM_PROMPT as EVALUATOR_PROMPT
from asgard.prompts.synthesizer import MAP_PROMPT, REDUCE_PROMPT, fill_repository

TOKEN = "ghp_bundle_secret"
_METRICS = {
    "approval_rate_iteration_1": 1.0,
    "approval_rate_final": 1.0,
    "mean_iterations": 1.0,
    "decision_flips": 0,
}


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()


def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _dump(path: Path, payload: object) -> None:
    _write(path, json.dumps(payload))


def _layout(root: Path) -> None:
    _dump(
        root / "config.json",
        {
            "repo": "acme/app",
            "github_token": TOKEN,
            "keys": {"openai": TOKEN},
            "llm": {"default": {"provider": "fake", "model": "class-model"}},
            "output_dir": str(root),
        },
    )
    _dump(root / "mine" / "output.json", [1])
    _dump(root / "mine" / "prs" / "1.json", {"body": TOKEN})
    _dump(root / "normalize" / "output.json", [{"context": {"pr_id": "1"}}])
    _write(root / "normalize" / "items.jsonl", "{}\n")
    _dump(
        root / "classify" / "output.json",
        [
            {
                "pr_id": "1",
                "status": "approved",
                "iterations": 1,
                "classification": {
                    "has_security_discussion": True,
                    "evidence": [{"quote": "bypass", "location": "description"}],
                    "rationale": "the description says bypass",
                },
                "evaluation": {"verdict": "APPROVE"},
            }
        ],
    )
    _write(root / "classify" / "items.jsonl", "{}\n")
    _write(
        root / "classify" / "trace.jsonl",
        "\n".join(
            [
                json.dumps({"input_tokens": 3, "output_tokens": 5}),
                json.dumps({"input_tokens": None, "output_tokens": None}),
                json.dumps({"role": "evaluator"}),
            ]
        )
        + "\n",
    )
    _dump(root / "classify" / "metrics.json", _METRICS)
    _dump(root / "consolidate" / "output.json", [{"pr_id": "1"}])
    _write(root / "consolidate" / "audit.jsonl", "{}\n")
    _dump(
        root / "consolidate" / "summary.json",
        {"total": 1, "security": 1, "non_security": 0, "unresolved": 0},
    )
    _dump(root / "synthesize" / "output.json", "Keep the bypass closed.")
    _dump(
        root / "security_md" / "output.json",
        {"path": "SECURITY.md", "unknown_citations": ["99"]},
    )
    _write(root / "SECURITY.md", "guideline\n")
    _write(
        root / "events.jsonl",
        "\n".join(
            [
                json.dumps(
                    {
                        "type": "StageStarted",
                        "stage": "mine",
                        "at": "2026-10-14T12:00:00+00:00",
                    }
                ),
                json.dumps(
                    {
                        "type": "StageFinished",
                        "stage": "mine",
                        "skipped": False,
                        "at": "2026-10-14T12:00:01+00:00",
                    }
                ),
                json.dumps({"type": "StageFinished", "stage": "mine", "skipped": True}),
                json.dumps({"type": "StageFinished", "stage": "classify", "skipped": False}),
                json.dumps({"type": "PRFiltered", "pr_id": "8"}),
            ]
        )
        + "\n",
    )
    _write(root / "notes.txt", "extra\n")


def _run(root: Path, *, export_xlsx: bool = False) -> dict[str, object]:
    stage = BundleStage(export_xlsx)
    manifest = stage.run(RunContext(root, lambda _event: None), None)
    assert isinstance(manifest, dict)
    return manifest


def test_synthesizer_prompt_hash_uses_override_files(tmp_path: Path) -> None:
    _layout(tmp_path)
    map_override = "Map override for {repository}."
    reduce_override = "Reduce override for {repository}."
    _write(tmp_path / "synthesize" / "map_prompt.txt", map_override)
    _write(tmp_path / "synthesize" / "reduce_prompt.txt", reduce_override)

    manifest = _run(tmp_path)

    assert manifest["prompts"]["synthesizer_map"] == _sha(
        fill_repository(map_override, "acme/app")
    )
    assert manifest["prompts"]["synthesizer_reduce"] == _sha(
        fill_repository(reduce_override, "acme/app")
    )
    assert manifest["prompts"]["synthesizer"] == manifest["prompts"]["synthesizer_reduce"]
    assert manifest["prompts"]["synthesizer_reduce"] != _sha(
        fill_repository(REDUCE_PROMPT, "acme/app")
    )


def test_different_classifier_versions_produce_different_manifest_hashes(
    tmp_path: Path,
) -> None:
    registry = PromptRegistry(tmp_path / "prompts")
    first = registry.save("classifier", name="alpha", body="classifier version alpha\n")
    second = registry.save("classifier", name="beta", body="classifier version beta\n")

    run_a = tmp_path / "run-a"
    run_b = tmp_path / "run-b"
    _layout(run_a)
    _layout(run_b)
    for run_dir, version_id in ((run_a, first.id), (run_b, second.id)):
        config = json.loads((run_dir / "config.json").read_text(encoding="utf-8"))
        config["prompts"] = {
            "classifier": version_id,
            "registry_dir": str(registry.root),
        }
        _dump(run_dir / "config.json", config)

    hash_a = _run(run_a)["prompts"]["classifier"]
    hash_b = _run(run_b)["prompts"]["classifier"]

    assert hash_a == _sha("classifier version alpha\n")
    assert hash_b == _sha("classifier version beta\n")
    assert hash_a != hash_b
    assert hash_a != _sha(CLASSIFIER_PROMPT)
    assert hash_b != _sha(CLASSIFIER_PROMPT)


def test_manifest_describes_every_file_and_omits_secrets(tmp_path: Path) -> None:
    _layout(tmp_path)

    manifest = _run(tmp_path)
    saved = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    copied = json.loads((tmp_path / "bundle" / "output.json").read_text(encoding="utf-8"))
    assert saved == manifest == copied

    on_disk = {
        path.relative_to(tmp_path).as_posix()
        for path in tmp_path.rglob("*")
        if path.is_file()
    }
    files = manifest["files"]
    assert isinstance(files, dict)
    assert on_disk == set(files)
    assert files["notes.txt"] == "unrecognized run file"
    assert files["mine/prs/1.json"] == "Cached GitHub pull request."
    assert files["manifest.json"] == "Run manifest."
    assert TOKEN not in json.dumps(manifest)
    assert "github_token" not in manifest["config"]
    assert "keys" not in manifest["config"]

    assert manifest["version"] == version("asgard")
    assert isinstance(manifest["commit"], str)
    assert len(manifest["commit"]) == 40
    prompts = manifest["prompts"]
    assert prompts == {
        "classifier": _sha(CLASSIFIER_PROMPT),
        "evaluator": _sha(EVALUATOR_PROMPT),
        "synthesizer_map": _sha(fill_repository(MAP_PROMPT, "acme/app")),
        "synthesizer_reduce": _sha(fill_repository(REDUCE_PROMPT, "acme/app")),
        "synthesizer": _sha(fill_repository(REDUCE_PROMPT, "acme/app")),
    }
    assert manifest["models"] == {
        "classifier": {"provider": "fake", "model": "class-model"},
        "evaluator": {"provider": "fake", "model": "class-model"},
        "synthesizer": {"provider": "fake", "model": "class-model"},
    }
    assert manifest["counts"] == {
        "mined": 1,
        "filtered": 1,
        "prefilter_in": None,
        "prefilter_out": None,
        "security": 1,
        "non_security": 0,
        "unresolved": 0,
    }
    assert manifest["loop"] == _METRICS
    assert manifest["citation_warnings"] == ["99"]
    assert manifest["tokens"] == {"input": 3, "output": 5}
    assert manifest["cost"] is None
    stages = manifest["stages"]
    assert stages["mine"] == {"items": 1, "duration_s": 1.0, "skipped": True}
    assert stages["classify"] == {"items": 1, "duration_s": None, "skipped": False}
    assert stages["normalize"] == {"items": 1, "duration_s": None, "skipped": None}
    assert stages["prefilter"] == {"items": None, "duration_s": None, "skipped": None}
    assert stages["bundle"]["items"] == 1
    assert stages["bundle"]["duration_s"] is None
    assert not (tmp_path / "bundle" / "classifications.xlsx").exists()


def test_prefilter_counts_come_from_removed_json(tmp_path: Path) -> None:
    _layout(tmp_path)
    _dump(tmp_path / "prefilter" / "output.json", [{"context": {"pr_id": "1"}}])
    _dump(
        tmp_path / "prefilter" / "removed.json",
        {
            "before": 3,
            "after": 1,
            "removed": [{"pr_id": "2", "reason": "no security keyword, label, or sensitive path"}],
        },
    )

    manifest = _run(tmp_path)

    assert manifest["counts"]["prefilter_in"] == 3
    assert manifest["counts"]["prefilter_out"] == 1
    assert manifest["stages"]["prefilter"]["items"] == 1
    files = manifest["files"]
    assert isinstance(files, dict)
    assert files["prefilter/output.json"] == "Pull requests kept for classification."
    assert files["prefilter/removed.json"] == "Pre-filter counts and the removed pull requests."


def test_export_xlsx_writes_one_row_per_classification(tmp_path: Path) -> None:
    _layout(tmp_path)

    manifest = _run(tmp_path, export_xlsx=True)

    path = tmp_path / "bundle" / "classifications.xlsx"
    assert path.is_file()
    assert TOKEN not in path.read_bytes().decode("utf-8", errors="ignore")
    assert manifest["files"]["bundle/classifications.xlsx"] == "Classification audit spreadsheet."
    sheet = load_workbook(path).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "pr_id",
        "status",
        "iterations",
        "decision",
        "rationale",
        "evidence",
        "verdict",
        "Evaluation",
        "Justify",
    )
    assert rows[1][0] == "1"
    assert rows[1][3] is True
    assert rows[1][6] == "APPROVE"
    assert rows[1][7] is None
    assert rows[1][8] is None


def test_export_xlsx_requires_openpyxl(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _layout(tmp_path)
    real_import = builtins.__import__

    def blocked(name: str, *args: object, **kwargs: object) -> object:
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("blocked")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked)
    with pytest.raises(ConfigError, match="openpyxl"):
        _run(tmp_path, export_xlsx=True)
    assert not (tmp_path / "bundle" / "classifications.xlsx").exists()


def test_export_xlsx_flag_is_stored(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("GITHUB_TOKEN", TOKEN)
    source = tmp_path / "asgard.toml"
    source.write_text(
        "\n".join(
            [
                'repo = "acme/app"',
                "export_xlsx = true",
                f'output_dir = "{tmp_path / "out"}"',
                "",
                "[llm.default]",
                'provider = "fake"',
                'model = "fake"',
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    config, _credentials = load_run(source)
    assert config.export_xlsx is True
